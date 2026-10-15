"""Adapted from adviSU pinned e1e2241 server/modules/export_utils.py; see SOURCE_LEDGER.json."""

import csv,io,re
from typing import Any

def _source_rows_to_csv(rows: list[dict[str, Any]], columns=None) -> bytes:
    if not rows and not columns:
        return b""
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=list(columns or rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")

def _source_rows_to_xlsx(rows: list[dict[str, Any]], *, sheet_name: str = "Analysis", columns=None) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("XLSX export requires openpyxl. Install server requirements first.") from exc

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]
    if rows or columns:
        headers = list(columns or rows[0])
        worksheet.append(headers)
        for row in rows:
            worksheet.append([row.get(header) for header in headers])
        header_fill = PatternFill("solid", fgColor="5265E0")
        for cell in worksheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column in worksheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
            worksheet.column_dimensions[column[0].column_letter].width = width

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def safe_cell(value):
    if not isinstance(value,str):return value
    value=re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]","",value)[:32000]
    # Excel/LibreOffice may ignore leading whitespace before a formula.
    if re.sub(r"^[\s\ufeff]*", "", value).startswith(("=","+","-","@")) or value.startswith(("\t","\r","\n")):value="'"+value
    return value

def safe_rows(rows,columns=None):
    if len(rows)>500:raise ValueError("Export en fazla 500 sonuç satırı.")
    columns=list(columns or (rows[0] if rows else []))
    headers=[str(safe_cell(k)) for k in columns]
    if len(set(headers))!=len(headers):raise ValueError("Güvenli export kolon adları çakışıyor; farklı AS adları verin.")
    return [{name:safe_cell(row.get(key)) for key,name in zip(columns,headers)} for row in rows],headers

def rows_to_csv(rows,columns=None):
    clean,headers=safe_rows(rows,columns)
    return _source_rows_to_csv(clean,headers)
def rows_to_xlsx(rows,columns=None):
    clean,headers=safe_rows(rows,columns)
    return _source_rows_to_xlsx(clean,columns=headers)
