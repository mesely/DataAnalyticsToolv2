from copy import deepcopy
import csv, io, json
from pathlib import Path
import pytest
from openpyxl import load_workbook
from analysis.worker import handle, run, METRICS, meta
from analysis.sql_lab import validate_local, compare_results
from analysis.adapted.exports import rows_to_csv, rows_to_xlsx
from analysis.adapted.identity import content_hash_text
from analysis.context import search_context


def sql(statement, filters=None, dataset="demo"):
    return handle(
        {
            "action": "sql",
            "dataset": dataset,
            "statement": statement,
            "filters": filters or {},
        }
    )


def test_sql_independent_oracle_and_scoped_view():
    oracle = json.loads(
        (Path(__file__).parents[1] / "data/demo_oracle.json").read_text()
    )
    result = sql(
        "SELECT SUM(line_pence) AS gross_pence, COUNT(DISTINCT invoice) AS orders, SUM(quantity) AS units FROM transactions"
    )
    assert result["rows"] == [
        {k: oracle[k] for k in ["gross_pence", "orders", "units"]}
    ]
    assert result["validation"]["verdict"] == "pass"
    assert sql(
        "SELECT SUM(line_pence) AS total FROM main.transactions", {"country": "France"}
    )["rows"] == [{"total": 703110}]
    assert sql("SELECT COUNT(*) AS count FROM transactions", {"start": "2030-01-01"})[
        "rows"
    ] == [{"count": 0}]
    assert sql(
        "WITH c AS (SELECT country, SUM(line_pence) AS p FROM transactions GROUP BY country) SELECT SUM(p) AS total FROM c"
    )["rows"] == [{"total": 2140500}]


@pytest.mark.parametrize(
    "statement",
    [
        "DROP TABLE transactions",
        "DELETE FROM transactions",
        "SELECT 1; SELECT 2",
        "SELECT * FROM read_csv_auto('/etc/passwd')",
        "SELECT read_text('/etc/passwd')",
        "SELECT * FROM read_parquet('https://evil.example/a')",
        "SELECT * FROM duckdb_tables()",
        "SELECT * FROM range(1000000000)",
        "SELECT * FROM transactions a CROSS JOIN transactions b",
        "WITH RECURSIVE x AS (SELECT 1 AS a UNION ALL SELECT a+1 FROM x) SELECT * FROM x",
        "SELECT customer_id FROM transactions",
        "SELECT unknown FROM transactions",
        "SELECT * FROM demo.main.transactions",
        "ATTACH '/tmp/other.duckdb' AS evil",
        "PRAGMA version",
        "INSTALL httpfs",
        "SELECT query('SELECT * FROM transactions')",
        "SELECT getenv('HOME')",
        "WITH d AS (DELETE FROM transactions RETURNING *) SELECT * FROM d",
    ],
)
def test_local_sql_refusal(statement):
    with pytest.raises(ValueError):
        validate_local(statement)


def test_limit_no_customer_and_unique_columns():
    result = sql("SELECT * FROM transactions ORDER BY invoice,product_id")
    assert result["truncated"] and len(result["rows"]) == 500
    assert "customer_id" not in result["columns"]
    with pytest.raises(ValueError):
        sql("SELECT 1 AS a, 2 AS a")


def test_sql_result_digest_and_source_identity():
    a = sql("SELECT SUM(line_pence) AS p FROM transactions")
    b = sql("SELECT SUM(line_pence) AS p FROM transactions")
    c = sql("SELECT SUM(line_pence) AS p FROM transactions", {"country": "France"})
    assert (
        a["evidence"]["output_sha256"]
        == b["evidence"]["output_sha256"]
        != c["evidence"]["output_sha256"]
    )
    assert a["evidence"]["source_id"].startswith("query:")


def test_sql_cannot_mutate_data_file():
    p = Path(__file__).parents[1] / ".local/demo.duckdb"
    before = content_hash_text(p.read_bytes().hex())
    sql("SELECT COUNT(*) AS c FROM transactions")
    assert before == content_hash_text(p.read_bytes().hex())


def test_bm25_normalized_scope_and_unknown():
    m = meta("demo")
    assert search_context("brut satis", m, METRICS)[0]["id"] == "gross_sales"
    assert search_context("sipariş", m, METRICS)[0]["id"] == "orders"
    assert search_context("line_pence", m, METRICS)[0]["id"] == "line_pence"
    assert search_context("zzzzzzunknown", m, METRICS) == []
    assert all(
        x["dataset_sha256"] == m["sha256"] for x in search_context("sales", m, METRICS)
    )


def test_csv_xlsx_formula_defense_and_numeric_types():
    rows = [
        {
            "=bad": ' =HYPERLINK("evil")',
            "product": "\t+SUM(1,2)",
            "amount": -4.25,
            "units": 3,
        }
    ]
    parsed = list(csv.DictReader(io.StringIO(rows_to_csv(rows).decode("utf-8-sig"))))[0]
    assert parsed["'=bad"].startswith("' ") and parsed["product"].startswith("'")
    book = load_workbook(io.BytesIO(rows_to_xlsx(rows)), data_only=False)
    sheet = book.active
    assert all(c.data_type != "f" for row in sheet for c in row)
    assert sheet["C2"].value == -4.25 and sheet["C2"].data_type == "n"
    with pytest.raises(ValueError):
        rows_to_csv([{"x": 1}] * 501)


def test_compatible_diff_scope_zero_and_mismatch():
    a = run({"dataset": "demo", "include_graph": False})
    b = run(
        {"dataset": "demo", "filters": {"country": "France"}, "include_graph": False}
    )
    diff = compare_results(a, b)
    assert diff["same_dataset_version"] and not diff["same_filters"]
    assert diff["metrics"][0]["delta"] == 703110 - 2140500
    z = deepcopy(a)
    z["summary"]["gross_pence"] = 0
    assert compare_results(z, b)["metrics"][0]["percent"] is None
    b["evidence"]["currency"] = "USD"
    with pytest.raises(ValueError):
        compare_results(a, b)


def test_full_artifact_digest_covers_graph(monkeypatch):
    from analysis import worker

    monkeypatch.setattr(
        worker, "graph", lambda r: {"nodes": [], "edges": [], "mode": "copurchase"}
    )
    a = run({"dataset": "demo"})
    monkeypatch.setattr(
        worker,
        "graph",
        lambda r: {"nodes": [{"id": "changed"}], "edges": [], "mode": "copurchase"},
    )
    b = run({"dataset": "demo"})
    assert a["evidence"]["recipe_version"] == "retail-analysis-v4"
    assert a["evidence"]["output_sha256"] != b["evidence"]["output_sha256"]


def test_exports_preserve_query_order_after_jsonb_and_empty_headers():
    rows = [{"orders": 60, "country": "France", "gross_sales_gbp": 7031.1}]
    columns = ["country", "gross_sales_gbp", "orders"]
    csv_rows = list(
        csv.reader(io.StringIO(rows_to_csv(rows, columns).decode("utf-8-sig")))
    )
    assert csv_rows[0] == columns and csv_rows[1] == ["France", "7031.1", "60"]
    assert rows_to_csv([], columns).decode("utf-8-sig").strip() == ",".join(columns)
    book = load_workbook(io.BytesIO(rows_to_xlsx([], columns)))
    assert list(book.active.values) == [tuple(columns)]
    with pytest.raises(ValueError):
        rows_to_csv([{"=x": 1, "'=x": 2}])


def test_spreadsheet_formula_after_bom_and_whitespace():
    row = {"product": ' \ufeff=HYPERLINK("evil")'}
    csv_rows = list(csv.DictReader(io.StringIO(rows_to_csv([row]).decode("utf-8-sig"))))
    assert csv_rows[0]["product"].startswith("'")
    sheet = load_workbook(io.BytesIO(rows_to_xlsx([row]))).active
    assert sheet["A2"].data_type != "f" and sheet["A2"].value.startswith("'")


@pytest.mark.parametrize(
    "question",
    [
        "What is net revenue?",
        "Why did sales drop?",
        "Predict next month revenue",
        "Net ciro nedir?",
    ],
)
def test_unavailable_question_semantics_abstain(question):
    r = run({"dataset": "demo", "question": question, "include_graph": False})
    assert r["answer"]["status"] in ["cannot_verify", "refused"]


def test_local_answer_numeric_claims_are_scoped_and_reject_tampering():
    from analysis.claims import verify_numeric_claims

    r = run(
        {
            "dataset": "demo",
            "question": "Which products sell the most?",
            "include_graph": False,
        }
    )
    answer = r["answer"]
    assert answer["claims"][0]["scope"]["product_id"] == r["products"][0]["product_id"]
    assert r["products"][0]["label"] in answer["text"]
    assert verify_numeric_claims(answer["claims"], r)
    wrong = deepcopy(answer["claims"])
    wrong[0]["value"] += 1
    assert not verify_numeric_claims(wrong, r)
    wrong = deepcopy(answer["claims"])
    wrong[0]["scope"]["product_id"] = "not-present"
    assert not verify_numeric_claims(wrong, r)
    assert not verify_numeric_claims([{"metric": "orders", "value": float("nan")}], r)
    assert not verify_numeric_claims([{"metric": "orders", "value": True}], r)
    assert not verify_numeric_claims([{"metric": "profit", "value": 10}], r)


@pytest.mark.parametrize(
    "question,reason",
    [
        ("Which products are bought together?", "open_network_tool"),
        ("Aylık satış trendi", "open_temporal_tool"),
    ],
)
def test_questions_requiring_other_tools_are_not_mislabeled_as_answered(
    question, reason
):
    r = run({"dataset": "demo", "question": question, "include_graph": False})
    assert r["answer"]["status"] == "clarify" and r["answer"]["reason"] == reason


def test_equal_rankings_have_deterministic_tie_breakers(monkeypatch, tmp_path):
    import pandas as pd
    from analysis import worker, prepare

    monkeypatch.setattr(worker, "DATA", tmp_path)
    monkeypatch.setattr(prepare, "DATA", tmp_path)
    rows = [
        ["I", "Z", "Z item", 2, 0.5, "2025-01-01", "Germany", "c"],
        ["I", "Z", "Z item", 1, 1.0, "2025-01-01", "France", "c"],
        ["J", "A", "A item", 1, 1.0, "2025-01-01", "France", "c"],
        ["J", "A", "A item", 2, 0.5, "2025-01-01", "Germany", "c"],
    ]
    dataset = "usr_abcdefabcdef_123456abcdef"
    prepare.store(
        dataset,
        "Tie fixture",
        pd.DataFrame(rows, columns=prepare.COLUMNS),
        {"kind": "synthetic", "publisher": "test", "license": "CC0"},
    )
    r = worker.run({"dataset": dataset, "include_graph": False})
    assert [x["product_id"] for x in r["products"]] == ["A", "Z"]
    assert [x["label"] for x in r["countries"]] == ["France", "Germany"]
    assert r["preview"][0]["quantity"] == 1
    assert (
        worker.run({"dataset": dataset, "include_graph": False})["evidence"][
            "output_sha256"
        ]
        == r["evidence"]["output_sha256"]
    )


def test_upload_rejects_price_precision_that_would_change_materialized_facts():
    text = "invoice,product_id,product,quantity,unit_price,order_date,country,customer_id\nI,P,Item,1000,0.12345,2025-01-01,France,C\n"
    with pytest.raises(ValueError, match="4 ondalık"):
        handle(
            {
                "action": "import",
                "csv": text,
                "owner": "abcdefabcdef",
                "title": "precision fixture",
            }
        )


def test_real_worker_accepts_valid_csv_with_large_json_quote_expansion():
    import subprocess, sys, uuid, hashlib
    from analysis.prepare import DATA, ROOT

    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(
        [
            "invoice",
            "product_id",
            "product",
            "quantity",
            "unit_price",
            "order_date",
            "country",
            "customer_id",
        ]
    )
    writer.writerow(["I", "P", '"' * 600000, 1, 1.0, "2025-01-01", "France", "C"])
    text = stream.getvalue()
    owner = uuid.uuid4().hex[:12]
    dataset = "usr_" + owner + "_" + hashlib.sha256(text.encode()).hexdigest()[:12]
    payload = json.dumps(
        {
            "action": "import",
            "csv": text,
            "owner": owner,
            "title": "escaped CSV fixture",
        }
    )
    assert len(text.encode()) < 2000000 and len(payload) > 2100000
    try:
        child = subprocess.run(
            [sys.executable, "-m", "analysis.worker"],
            input=payload,
            text=True,
            capture_output=True,
            cwd=ROOT.parent,
            timeout=20,
        )
        assert child.returncode == 0
        response = json.loads(child.stdout)
        assert response["ok"] and response["result"]["rows"] == 1
    finally:
        # Only this test's randomly namespaced disposable files are removed.
        for suffix in [".json", ".duckdb"]:
            (DATA / (dataset + suffix)).unlink(missing_ok=True)
