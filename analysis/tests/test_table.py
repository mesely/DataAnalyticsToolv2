"""v3 G3: the overview's data table — correctness against an independent oracle, paging, filters, downloads, safety."""

import base64
import csv
import io
import json
from collections import defaultdict
from decimal import Decimal

import pandas as pd
import pytest

from analysis import prepare, worker
from analysis.prepare import ROOT
from analysis.table import MAX_EXPORT_CELLS, ask, export, spec_to_table, table

DATASET = "usr_7ab1e07ab1e0_7ab1e07ab1e0"
ROWS = [
    ["1001", "A1", "MUG", 10, 2.0, "2025-01-03", "Germany", "c1"],
    ["1001", "B2", "=HYPERLINK(\"x\")", 2, 5.0, "2025-01-03", "Germany", "c1"],
    ["1002", "A1", "MUG", 1, 2.0, "2025-02-10", "France", "c2"],
    ["1002", "B2", "=HYPERLINK(\"x\")", 1, 5.0, "2025-02-10", "France", "c2"],
    ["1003", "C3", "PLATE", 3, 1.5, "2025-02-11", "France", "c3"],
    ["1004", "POST", "POSTAGE", 1, 18.0, "2025-02-12", "Germany", None],
]


@pytest.fixture()
def ds(monkeypatch, tmp_path):
    monkeypatch.setattr(worker, "DATA", tmp_path)
    monkeypatch.setattr(prepare, "DATA", tmp_path)
    monkeypatch.setattr("analysis.chart_studio.DATA", tmp_path)
    prepare.store(DATASET, "Table fixture", pd.DataFrame(ROWS, columns=prepare.COLUMNS),
                  {"kind": "synthetic", "publisher": "test", "license": "CC0"})
    return DATASET


def test_rows_page_and_total_order(ds):
    r = table({"dataset": ds, "view": "rows", "page_size": 25})
    assert r["total"] == 6 and r["pages"] == 1
    assert [c["name"] for c in r["columns"]] == ["invoice", "order_date", "country", "product_id", "product", "quantity", "unit_price", "line_pence"]
    assert "customer_id" not in r["rows"][0]
    assert r["rows"][0]["unit_price"] == "2" and r["rows"][2]["line_pence"] == 200
    # Sorting by price is numeric, not text ("18" after "5").
    s = table({"dataset": ds, "view": "rows", "page_size": 25, "sort": [{"column": "unit_price", "dir": "desc"}]})
    assert [Decimal(x["unit_price"]) for x in s["rows"]] == sorted((Decimal(str(x[4])) for x in ROWS), reverse=True)


def test_summary_matches_an_independent_oracle(ds):
    oracle = defaultdict(lambda: {"pence": 0, "orders": set(), "units": 0})
    for inv, _, _, q, p, _, country, _ in ROWS:
        o = oracle[country]
        o["pence"] += int(Decimal(str(q)) * Decimal(str(p)) * 100)
        o["orders"].add(inv)
        o["units"] += q
    r = table({"dataset": ds, "view": "summary", "group_by": ["country"], "page_size": 25})
    got = {x["country"]: x for x in r["rows"]}
    for country, o in oracle.items():
        assert got[country]["gross_pence"] == o["pence"]
        assert got[country]["orders"] == len(o["orders"]) and got[country]["units"] == o["units"]
        assert got[country]["aov_pence"] == round(o["pence"] / len(o["orders"]))
    assert [x["country"] for x in r["rows"]] == ["Germany", "France"]  # default: gross sales, largest first


def test_product_summary_names_its_columns_apart_and_skips_service_lines(ds):
    r = table({"dataset": ds, "view": "summary", "group_by": ["product"], "page_size": 25})
    names = [c["name"] for c in r["columns"]]
    assert names[:2] == ["product_id", "product"] and len(names) == len(set(names))
    assert "POST" not in {x["product_id"] for x in r["rows"]}
    both = table({"dataset": ds, "view": "summary", "group_by": ["product"], "page_size": 25, "include_service": True})
    assert "POST" in {x["product_id"] for x in both["rows"]}


def test_column_filters(ds):
    def total(filters, view="rows", groups=None):
        req = {"dataset": ds, "view": view, "page_size": 25, "filters_table": filters}
        if groups:
            req["group_by"] = groups
        return table(req)["total"]

    assert total([{"column": "product", "op": "contains", "value": "mu"}]) == 2
    assert total([{"column": "product", "op": "contains", "value": "%"}]) == 0  # LIKE wildcards are literal
    assert total([{"column": "quantity", "op": "gte", "value": 2}, {"column": "quantity", "op": "lte", "value": 3}]) == 2
    assert total([{"column": "order_date", "op": "gte", "value": "2025-02-01"}]) == 4
    assert total([{"column": "country", "op": "in", "value": ["France"]}]) == 3
    assert total([{"column": "gross_pence", "op": "gte", "value": 1000}], "summary", ["country"]) == 2
    assert total([{"column": "orders", "op": "gte", "value": 3}], "summary", ["country"]) == 0


@pytest.mark.parametrize("bad", [
    {"view": "rows", "filters_table": [{"column": "customer_id", "op": "eq", "value": "c1"}]},
    {"view": "rows", "filters_table": [{"column": "product); DROP TABLE transactions;--", "op": "eq", "value": "x"}]},
    {"view": "rows", "filters_table": [{"column": "product", "op": "like", "value": "x"}]},
    {"view": "rows", "filters_table": [{"column": "quantity", "op": "contains", "value": "1"}]},
    {"view": "rows", "sort": [{"column": "line_pence; SELECT 1", "dir": "asc"}]},
    {"view": "rows", "sort": [{"column": "invoice", "dir": "sideways"}]},
    {"view": "summary", "group_by": ["customer_id"]},
    {"view": "summary", "group_by": ["country", "month", "week"]},
    {"view": "rows", "page_size": 1000},
    {"view": "rows", "page": 0},
])
def test_unsafe_or_unknown_requests_are_refused(ds, bad):
    with pytest.raises(ValueError):
        table({"dataset": ds, **bad})


def test_exports_neutralise_formulas_and_match_the_screen(ds):
    e = export({"dataset": ds, "view": "rows", "format": "csv", "columns": ["invoice", "product", "line_pence"],
                "labels": {"line_pence": "Tutar"}, "sort": [{"column": "invoice", "dir": "asc"}]})
    text = base64.b64decode(e["base64"]).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0] == ["Fatura", "Ürün", "Tutar"]
    assert rows[2][1].startswith("'=")  # the spreadsheet will not evaluate it
    assert len(rows) == 7 and e["rows"] == 6 and e["columns"] == 3
    x = export({"dataset": ds, "view": "summary", "group_by": ["country"], "format": "xlsx", "locale": "en"})
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(base64.b64decode(x["base64"]))).active
    values = [[c.value for c in r] for r in ws.iter_rows()]
    assert values[0][:2] == ["Country", "Gross sales"] and values[1][0] == "Germany" and values[1][1] == 48.0  # postage is in gross sales


def test_large_downloads_are_refused_with_a_way_forward(monkeypatch, ds):
    monkeypatch.setattr("analysis.table.MAX_EXPORT_CELLS", 10)
    with pytest.raises(ValueError, match="summary|özet"):
        export({"dataset": ds, "view": "rows", "format": "csv"})
    assert MAX_EXPORT_CELLS == 500_000


def test_ask_turns_a_chart_plan_into_a_summary(ds):
    r = ask({"dataset": ds, "prompt": "top 3 products by units", "locale": "en"})
    assert r["view"] == "summary" and r["group_by"] == ["product"]
    assert r["asked"]["planner"]["routing"]["task"] == "chart_plan"
    assert r["rows"][0]["product"] == "MUG" and len(r["rows"]) <= 3
    t = spec_to_table({"metric": "orders", "dimension": "month", "countries": ["Germany", "France"], "split_by_country": True,
                       "exclude_countries": [], "top_n": 10, "sort": "natural", "start": "2025-01-01", "end": None})
    assert t["group_by"] == ["month", "country"] and t["sort"] == [{"column": "month", "dir": "asc"}, {"column": "country", "dir": "asc"}]
    assert t["filters_table"] == [{"column": "country", "op": "in", "value": ["Germany", "France"]}] and t["dates"] == {"start": "2025-01-01"}


def test_dispatch_through_the_worker(ds):
    assert worker.handle({"action": "table", "dataset": ds, "view": "rows", "page_size": 25})["total"] == 6
